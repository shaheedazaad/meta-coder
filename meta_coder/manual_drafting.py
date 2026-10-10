"""Draft a coding manual from an uploaded PDF, DOCX, RTF, or Markdown document.

This module owns the two trust boundaries in that workflow: extracting bounded
plain text from an untrusted upload, and validating the model's structured draft
through the same ``CodingManual`` validator used by the editor and YAML import.
It deliberately does not write project files; a draft must first be reviewed in
the structured editor and explicitly saved by the user.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
import re
from typing import Any, BinaryIO

from docx import Document
from pypdf import PdfReader
from striprtf.striprtf import rtf_to_text

from .extraction import ProviderError, parse_json_response
from .manual import CodingManual, ManualError, manual_from_editor_payload


SUPPORTED_DOCUMENT_SUFFIXES = {".pdf", ".docx", ".rtf", ".md"}
MAX_EXTRACTED_TEXT_CHARS = 500_000


class ManualDraftError(ValueError):
    """Raised when an uploaded manual cannot produce a reviewable draft."""


def _read_bounded(stream: BinaryIO, *, max_bytes: int) -> bytes:
    chunks: list[bytes] = []
    total = 0
    while True:
        chunk = stream.read(1024 * 1024)
        if not chunk:
            break
        total += len(chunk)
        if total > max_bytes:
            raise ManualDraftError(
                f"The manual document exceeds the {max_bytes // (1024 * 1024)} MB upload limit."
            )
        chunks.append(chunk)
    if not chunks:
        raise ManualDraftError("The manual document is empty.")
    return b"".join(chunks)


def pdf_text(data: bytes) -> str:
    if not data.startswith(b"%PDF"):
        raise ManualDraftError("The uploaded file is not a valid PDF.")
    try:
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ManualDraftError("The PDF is password-protected and cannot be read.")
        pages = []
        for number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                pages.append(f"[Page {number}]\n{text}")
    except ManualDraftError:
        raise
    except Exception as exc:
        raise ManualDraftError(f"The PDF could not be read: {exc}") from exc
    return "\n\n".join(pages)


def _docx_text(data: bytes) -> str:
    # DOCX is a ZIP container. Checking the signature gives a clearer error than
    # passing arbitrary bytes into python-docx and exposing its implementation
    # exception to the user.
    if not data.startswith(b"PK"):
        raise ManualDraftError("The uploaded file is not a valid DOCX document.")
    try:
        document = Document(BytesIO(data))
        parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append("\t".join(cells))
    except Exception as exc:
        raise ManualDraftError(f"The DOCX document could not be read: {exc}") from exc
    return "\n".join(parts)


def _rtf_text(data: bytes) -> str:
    if not data.startswith(b"{\\rtf"):
        raise ManualDraftError("The uploaded file is not a valid RTF document.")
    try:
        # Express raw non-ASCII bytes as RTF escapes so the parser decodes them
        # using the document's code page, just like existing hex escapes.
        escaped = re.sub(rb"[\x80-\xff]", lambda match: f"\\'{match[0][0]:02x}".encode("ascii"), data)
        return rtf_to_text(escaped.decode("ascii"))
    except Exception as exc:
        raise ManualDraftError(f"The RTF document could not be read: {exc}") from exc


def _markdown_text(data: bytes) -> str:
    try:
        # Keep headings, lists, and tables as written; the model reads Markdown.
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ManualDraftError("The Markdown document must be saved as UTF-8 text.") from exc
    if "\x00" in text:
        raise ManualDraftError("The uploaded file is not a valid Markdown text document.")
    return text


def extract_manual_document_text(
    stream: BinaryIO, filename: str, *, max_bytes: int
) -> str:
    """Validate an uploaded manual and return bounded, non-empty text."""

    suffix = Path(filename.replace("\\", "/")).suffix.lower()
    if suffix not in SUPPORTED_DOCUMENT_SUFFIXES:
        raise ManualDraftError("Upload a PDF, DOCX, RTF, or Markdown (.md) coding-manual document.")
    data = _read_bounded(stream, max_bytes=max_bytes)
    extractors = {".pdf": pdf_text, ".docx": _docx_text, ".rtf": _rtf_text, ".md": _markdown_text}
    text = extractors[suffix](data)
    text = text.strip()
    if not text:
        raise ManualDraftError(
            "No readable text was found in the document. Scanned PDFs need OCR before import."
        )
    if len(text) > MAX_EXTRACTED_TEXT_CHARS:
        raise ManualDraftError(
            "The extracted manual text is too long to draft safely. Split the manual into a "
            "smaller document and try again."
        )
    return text


DRAFT_PROMPT = """You convert an existing human-written meta-analysis coding manual into
MetaCoder's structured coding-manual schema.

Create a faithful, reviewable draft from the document below. Do not invent fields,
category levels, or instructions that the document does not support.

Schema rules:
- name: a short snake_case name for this manual.
- description: a concise description; use an empty string if the document has none.
- effect_definition: the document's definition of the effect/comparison/unit of
  analysis. It must be specific enough to distinguish multiple effects in one paper.
- effects: fields whose values should be coded from research articles. Do not include
  administrative/source-identification fields such as row ID, filename, citation,
  authors, year, effect locator, reviewer notes, or page number.
- Never name an effect row_id, source_pdf, locator, authors, year, or status (in
  any capitalization); MetaCoder reserves these. If the document codes such a
  property of the study, use a more specific name, e.g. publication_status.
- Each effect has a unique name, one type (string, number, integer, or boolean), a
  self-contained coding instruction in description, evidence_required, and levels.
- Use levels only for categorical string fields. Preserve every documented category
  label and its meaning. Otherwise return an empty levels array.
- Numbered categories (e.g. 1 = spatial, 2 = temporal) are string fields with
  string labels "1" and "2", not integer fields. Any field with a non-empty levels
  array must have type "string". Numeric measurements and booleans use levels: [].
- Set evidence_required true unless the field is explicitly not sourced from the
  research article.
- Do not add a notes field; MetaCoder adds its standard notes field automatically.

Return only the schema-constrained object. The user will review it before saving.

CODING MANUAL DOCUMENT:
"""


def build_manual_draft_prompt(document_text: str) -> str:
    return DRAFT_PROMPT + document_text


def build_manual_draft_schema(*, dialect: str = "gemini") -> dict[str, Any]:
    """Structured-output schema for the editor's list-based manual payload."""

    if dialect not in {"gemini", "json_schema"}:
        raise ValueError("Unknown manual-draft schema dialect.")
    upper = dialect == "gemini"
    types = {
        "object": "OBJECT" if upper else "object",
        "array": "ARRAY" if upper else "array",
        "string": "STRING" if upper else "string",
        "boolean": "BOOLEAN" if upper else "boolean",
    }

    def closed_object(properties: dict[str, Any], required: list[str]) -> dict[str, Any]:
        schema: dict[str, Any] = {
            "type": types["object"],
            "properties": properties,
            "required": required,
        }
        if not upper:
            schema["additionalProperties"] = False
        return schema

    level = closed_object(
        {
            "value": {"type": types["string"]},
            "description": {"type": types["string"]},
        },
        ["value", "description"],
    )
    effect = closed_object(
        {
            "name": {"type": types["string"]},
            "type": {
                "type": types["string"],
                "enum": ["string", "number", "integer", "boolean"],
            },
            "description": {"type": types["string"]},
            "evidence_required": {"type": types["boolean"]},
            "levels": {"type": types["array"], "items": level},
        },
        ["name", "type", "description", "evidence_required", "levels"],
    )
    return closed_object(
        {
            "name": {"type": types["string"]},
            "description": {"type": types["string"]},
            "effect_definition": {"type": types["string"]},
            "effects": {"type": types["array"], "items": effect},
        },
        ["name", "description", "effect_definition", "effects"],
    )


def parse_manual_draft_response(raw_text: str) -> CodingManual:
    """Parse model JSON and enforce the canonical manual invariants."""

    try:
        payload, _repaired = parse_json_response(raw_text)
        # Models sometimes call numbered categories integers (or yes/no
        # categories booleans). Categories are strings in our editor/schema.
        # Preserve their labels and instructions in the reviewable draft instead
        # of dropping levels to make the model's suggested type validate.
        effects = payload.get("effects") if isinstance(payload, dict) else None
        if isinstance(effects, list):
            for field in effects:
                if (
                    isinstance(field, dict)
                    and str(field.get("type", "")).strip().lower() in {"number", "integer", "boolean"}
                    and isinstance(field.get("levels"), list)
                    and field["levels"]
                ):
                    field["type"] = "string"
        return manual_from_editor_payload(payload)
    except (ProviderError, ManualError) as exc:
        raise ManualDraftError(f"The model returned an invalid coding-manual draft: {exc}") from exc
