from io import BytesIO
import json

import pytest
from docx import Document

import meta_coder.manual_drafting as manual_drafting
from meta_coder.manual_drafting import (
    ManualDraftError,
    build_manual_draft_prompt,
    build_manual_draft_schema,
    extract_manual_document_text,
    parse_manual_draft_response,
)


def test_extracts_text_and_tables_from_docx_manual():
    stream = BytesIO()
    document = Document()
    document.add_heading("Affordance Meta-analysis Coding Manual")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Age Mean"
    table.cell(0, 1).text = "Mean age of participants"
    document.save(stream)
    stream.seek(0)

    text = extract_manual_document_text(stream, "manual.docx", max_bytes=1024 * 1024)

    assert "Affordance Meta-analysis Coding Manual" in text
    assert "Age Mean" in text
    assert "Mean age of participants" in text


def test_extracts_page_marked_text_from_pdf(monkeypatch):
    class Page:
        def extract_text(self):
            return "ORIGINAL RESEARCH"

    class Reader:
        is_encrypted = False
        pages = [Page()]

    monkeypatch.setattr(manual_drafting, "PdfReader", lambda _stream: Reader())
    text = extract_manual_document_text(BytesIO(b"%PDF-synthetic"), "manual.pdf", max_bytes=1024)

    assert text.startswith("[Page 1]")
    assert "ORIGINAL RESEARCH" in text


def test_rejects_wrong_document_type_and_spoofed_docx():
    with pytest.raises(ManualDraftError, match="PDF, DOCX, RTF, or Markdown"):
        extract_manual_document_text(BytesIO(b"text"), "manual.txt", max_bytes=100)
    with pytest.raises(ManualDraftError, match="not a valid DOCX"):
        extract_manual_document_text(BytesIO(b"not a zip"), "manual.docx", max_bytes=100)


def test_rejects_document_over_upload_limit_before_parsing():
    with pytest.raises(ManualDraftError, match="upload limit"):
        extract_manual_document_text(BytesIO(b"PK" + b"x" * 20), "manual.docx", max_bytes=10)


def test_extracts_rtf_formatting_unicode_and_table_cells():
    data = br"{\rtf1\ansi\ansicpg1252{\fonttbl{\f0 Arial;}}\b Coding manual\b0\par Age\cell Mean age\cell\row Caf\'e9 \u945?}"
    text = extract_manual_document_text(BytesIO(data), "manual.RTF", max_bytes=1024)
    assert "Coding manual" in text
    assert "Age" in text and "Mean age" in text
    assert "Café α" in text
    assert "Arial" not in text and "\\" not in text


def test_rtf_raw_bytes_follow_document_codepage():
    data = b"{\\rtf1\\ansi\\ansicpg1251 " + "Возраст".encode("cp1251") + b"}"
    assert extract_manual_document_text(BytesIO(data), "manual.rtf", max_bytes=1024) == "Возраст"


def test_markdown_preserves_structure_and_accepts_utf8_bom():
    text = "# Coding manual\n\n- Effect: Δ response time\n\n| Field | Description |\n| --- | --- |\n| Age | Mean age |"
    assert extract_manual_document_text(
        BytesIO(text.encode("utf-8-sig")), "manual.MD", max_bytes=1024
    ) == text


@pytest.mark.parametrize("filename,data,message", [
    ("manual.rtf", b"not rtf", "not a valid RTF"),
    ("manual.rtf", br"{\rtf1\ansi }", "No readable text"),
    ("manual.md", b" \n\t", "No readable text"),
    ("manual.md", b"\xff", "UTF-8"),
    ("manual.md", b"binary\x00data", "not a valid Markdown"),
])
def test_rejects_unreadable_rtf_and_markdown(filename, data, message):
    with pytest.raises(ManualDraftError, match=message):
        extract_manual_document_text(BytesIO(data), filename, max_bytes=1024)


@pytest.mark.parametrize("filename,data", [
    ("manual.rtf", br"{\rtf1\ansi Coding manual}"),
    ("manual.md", b"# Coding manual"),
])
def test_new_formats_obey_upload_and_extracted_text_limits(monkeypatch, filename, data):
    with pytest.raises(ManualDraftError, match="upload limit"):
        extract_manual_document_text(BytesIO(data), filename, max_bytes=5)
    monkeypatch.setattr(manual_drafting, "MAX_EXTRACTED_TEXT_CHARS", 5)
    with pytest.raises(ManualDraftError, match="too long"):
        extract_manual_document_text(BytesIO(data), filename, max_bytes=1024)


def test_draft_response_uses_canonical_manual_validation_and_adds_notes():
    manual = parse_manual_draft_response(
        """{
          "name": "attention_review",
          "description": "Attention studies",
          "effect_definition": "Difference between cued and uncued trials",
          "effects": [{
            "name": "cue_type",
            "type": "string",
            "description": "Type of cue",
            "evidence_required": true,
            "levels": [{"value": "valid", "description": "Cue matches target"}]
          }]
        }"""
    )

    assert manual.effect_definition == "Difference between cued and uncued trials"
    assert set(manual.effects) == {"cue_type", "notes"}


def test_invalid_model_draft_is_not_accepted():
    with pytest.raises(ManualDraftError, match="invalid coding-manual draft"):
        parse_manual_draft_response(
            '{"name":"bad","description":"","effect_definition":"","effects":[]}'
        )


@pytest.mark.parametrize("field_type", ["integer", "number", "boolean"])
def test_draft_preserves_categories_when_model_uses_non_string_type(field_type):
    payload = {
        "name": "dimension_review",
        "description": "",
        "effect_definition": "Difference between conditions",
        "effects": [{
            "name": "dimension",
            "type": field_type,
            "description": "Code the dimension using the documented categories",
            "evidence_required": True,
            "levels": [
                {"value": "1", "description": "Spatial"},
                {"value": "2", "description": "Temporal"},
            ],
        }],
    }
    manual = parse_manual_draft_response(json.dumps(payload))
    field = manual.effects["dimension"]
    assert field.type == "string"
    assert [(level.value, level.description) for level in field.levels] == [
        ("1", "Spatial"), ("2", "Temporal"),
    ]
    assert field.description == payload["effects"][0]["description"]
    assert field.evidence_required is True


def test_schema_dialects_match_provider_requirements():
    gemini = build_manual_draft_schema()
    json_schema = build_manual_draft_schema(dialect="json_schema")

    assert gemini["type"] == "OBJECT"
    assert json_schema["type"] == "object"
    assert json_schema["additionalProperties"] is False
    assert (
        json_schema["properties"]["effects"]["items"]["properties"]["type"]["enum"]
        == ["string", "number", "integer", "boolean"]
    )


@pytest.mark.parametrize("field_type", ["integer", "number", "boolean"])
def test_draft_keeps_non_categorical_field_types(field_type):
    payload = {
        "effect_definition": "Difference between conditions",
        "effects": [{"name": "measurement", "type": field_type, "levels": []}],
    }
    manual = parse_manual_draft_response(json.dumps(payload))
    assert manual.effects["measurement"].type == field_type


@pytest.mark.parametrize("levels, message", [
    ([{"value": "1"}, {"value": "1"}], "duplicate value"),
    ([{"value": ""}], "cannot be empty"),
    (["category"], "must be a mapping"),
    ("categories", "must be a list"),
])
def test_draft_still_rejects_invalid_category_definitions(levels, message):
    payload = {
        "effect_definition": "Difference between conditions",
        "effects": [{"name": "dimension", "type": "integer", "levels": levels}],
    }
    with pytest.raises(ManualDraftError, match=message):
        parse_manual_draft_response(json.dumps(payload))


def test_prompt_marks_document_and_requires_reviewable_scope():
    prompt = build_manual_draft_prompt("FIELD DEFINITIONS")
    assert prompt.endswith("FIELD DEFINITIONS")
    assert "Do not invent fields" in prompt
    assert "Do not add a notes field" in prompt


def test_prompt_and_draft_validation_reject_reserved_field_names():
    from meta_coder.manual import BASE_COLUMNS

    prompt = build_manual_draft_prompt("")
    assert all(name in prompt for name in BASE_COLUMNS)
    assert "publication_status" in prompt
    draft = {
        "name": "m", "description": "", "effect_definition": "comparison",
        "effects": [{"name": "Status", "type": "string", "description": "",
                     "evidence_required": True, "levels": []}],
    }
    with pytest.raises(ManualDraftError, match="`Status` is reserved.*publication_status"):
        parse_manual_draft_response(json.dumps(draft))


def test_empty_upload_and_invalid_pdf_signature():
    import io
    from meta_coder import manual_drafting as drafting
    with pytest.raises(drafting.ManualDraftError, match='empty'):
        drafting.extract_manual_document_text(io.BytesIO(b''), 'manual.md', max_bytes=100)
    with pytest.raises(drafting.ManualDraftError, match='not a valid PDF'):
        drafting.pdf_text(b'not pdf')


def test_password_protected_pdf_is_rejected():
    import io
    from pypdf import PdfWriter
    from meta_coder import manual_drafting as drafting
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt('secret')
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(drafting.ManualDraftError, match='password-protected'):
        drafting.pdf_text(buffer.getvalue())


def test_blank_pdf_pages_are_skipped():
    import io
    from pypdf import PdfWriter
    from meta_coder import manual_drafting as drafting
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = io.BytesIO()
    writer.write(buffer)
    assert drafting.pdf_text(buffer.getvalue()) == ''


def test_corrupt_docx_and_rtf_decoder_errors_are_wrapped(monkeypatch):
    from meta_coder import manual_drafting as drafting
    with pytest.raises(drafting.ManualDraftError, match='DOCX document could not be read'):
        drafting._docx_text(b'PKbroken')
    def fail(_):
        raise ValueError('bad RTF')
    monkeypatch.setattr(drafting, 'rtf_to_text', fail)
    with pytest.raises(drafting.ManualDraftError, match='RTF document could not be read'):
        drafting._rtf_text(b'{\\rtf1 broken}')


def test_empty_docx_tables_do_not_create_content():
    import io
    from docx import Document
    from meta_coder import manual_drafting as drafting
    document = Document()
    document.add_table(rows=1, cols=2)
    buffer = io.BytesIO()
    document.save(buffer)
    assert drafting._docx_text(buffer.getvalue()) == ''


def test_unknown_draft_schema_and_nonobject_payload():
    from meta_coder import manual_drafting as drafting
    with pytest.raises(ValueError, match='Unknown'):
        drafting.build_manual_draft_schema(dialect='xml')
    with pytest.raises(drafting.ManualDraftError, match='invalid coding-manual draft'):
        drafting.parse_manual_draft_response('[]')
